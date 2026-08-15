from pathlib import Path
import os
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document
import json

from fs_tools import read_file


def make_pdf(path: Path, text: str):
    c = canvas.Canvas(str(path), pagesize=letter)
    width, height = letter
    y = height - 72
    for line in text.splitlines():
        c.drawString(72, y, line)
        y -= 14
        if y < 72:
            c.showPage()
            y = height - 72
    c.save()


def make_docx(path: Path, text: str):
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    # add a simple table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Level"
    table.cell(1, 0).text = "Python"
    table.cell(1, 1).text = "Expert"
    doc.save(str(path))


def run():
    outdir = Path(os.environ.get("RESUMES_DIR", "resumes"))
    outdir.mkdir(exist_ok=True)
    text = "John Example\nSenior Software Engineer\nExperienced in Python, ML, and backend systems."
    pdf_path = outdir / "resume_john_example.pdf"
    docx_path = outdir / "resume_jane_example.docx"
    make_pdf(pdf_path, text)
    make_docx(docx_path, text)

    pdf_res = read_file(str(pdf_path))
    docx_res = read_file(str(docx_path))

    print("PDF read result:\n", json.dumps({k: v for k, v in pdf_res.items() if k != 'content'}, indent=2))
    print("DOCX read result:\n", json.dumps({k: v for k, v in docx_res.items() if k != 'content'}, indent=2))
    print("PDF content excerpt:\n", pdf_res.get('content','')[:400])
    print("DOCX content excerpt:\n", docx_res.get('content','')[:400])


if __name__ == '__main__':
    run()
