import os
import re
from pathlib import Path
from datetime import datetime
from typing import List, Dict, Any, Optional

try:
    import PyPDF2
except Exception:
    PyPDF2 = None

# optional better PDF extractors
try:
    import pdfplumber
except Exception:
    pdfplumber = None

try:
    from pdfminer.high_level import extract_text as pdfminer_extract_text
except Exception:
    pdfminer_extract_text = None

try:
    import docx
except Exception:
    docx = None


def _get_metadata(filepath: str) -> Dict[str, Any]:
    p = Path(filepath)
    stat = p.stat()
    return {
        "name": p.name,
        "path": str(p.resolve()),
        "size": stat.st_size,
        "modified": datetime.fromtimestamp(stat.st_mtime).isoformat(),
    }


def read_file(filepath: str) -> Dict[str, Any]:
    """Read resume files (PDF, TXT, DOCX) and return content + metadata.

    Returns dict: {status: 'success'|'error', content: str, metadata: {...}, error?: str}
    """
    p = Path(filepath)
    if not p.exists():
        return {"status": "error", "error": "file not found", "path": str(p)}

    suffix = p.suffix.lower()
    try:
        if suffix == ".pdf":
            # Try multiple extractors in order of preference
            text = ""
            # 1) pdfplumber (if available)
            if pdfplumber is not None:
                try:
                    with pdfplumber.open(str(p)) as pdf:
                        pages = [page.extract_text() or "" for page in pdf.pages]
                        text = "\n".join(pages).strip()
                except Exception:
                    text = ""

            # 2) pdfminer
            if not text and pdfminer_extract_text is not None:
                try:
                    text = pdfminer_extract_text(str(p)) or ""
                except Exception:
                    text = ""

            # 3) PyPDF2 as fallback
            if not text and PyPDF2 is not None:
                try:
                    text_parts: List[str] = []
                    with open(p, "rb") as fh:
                        reader = PyPDF2.PdfReader(fh)
                        for page in reader.pages:
                            try:
                                text_parts.append(page.extract_text() or "")
                            except Exception:
                                pass
                    text = "\n".join(text_parts)
                except Exception:
                    text = ""

            if not text:
                return {"status": "error", "error": "no PDF text extractor available or extraction failed (install pdfplumber/pdfminer.six/PyPDF2)"}
        elif suffix in (".docx", ".doc"):
            if docx is None:
                return {"status": "error", "error": "python-docx not installed. Install requirements."}
            try:
                doc = docx.Document(str(p))
                paragraphs = [para.text for para in doc.paragraphs if para.text]
                # extract tables as additional lines
                tables_text = []
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells)
                        tables_text.append(row_text)
                parts = []
                if paragraphs:
                    parts.append("\n".join(paragraphs))
                if tables_text:
                    parts.append("\n".join(tables_text))
                text = "\n\n".join(parts)
            except Exception as e:
                return {"status": "error", "error": f"docx extraction failed: {e}"}
        else:
            # treat as text
            with open(p, "r", encoding="utf-8", errors="ignore") as fh:
                text = fh.read()

        parsed = _parse_resume_text(text)
        return {"status": "success", "content": text, "metadata": _get_metadata(str(p)), "parsed": parsed}
    except Exception as e:
        return {"status": "error", "error": str(e)}


def list_files(directory: str, extension: Optional[str] = None) -> List[Dict[str, Any]]:
    """List files in a directory with optional extension filter. Returns list of metadata dicts."""
    p = Path(directory)
    if not p.exists() or not p.is_dir():
        return []
    ext = None
    if extension:
        ext = extension.lower()
        if not ext.startswith("."):
            ext = "." + ext

    results: List[Dict[str, Any]] = []
    for child in p.iterdir():
        if child.is_file():
            if ext and child.suffix.lower() != ext:
                continue
            meta = _get_metadata(str(child))
            results.append(meta)
    return results


def write_file(filepath: str, content: str) -> Dict[str, Any]:
    """Write content to a file, creating directories if needed."""
    p = Path(filepath)
    try:
        if not p.parent.exists():
            p.parent.mkdir(parents=True, exist_ok=True)
        with open(p, "w", encoding="utf-8") as fh:
            fh.write(content)
        return {"status": "success", "metadata": _get_metadata(str(p))}
    except Exception as e:
        return {"status": "error", "error": str(e)}


def search_in_file(filepath: str, keyword: str, context: int = 100) -> Dict[str, Any]:
    """Search for keyword in file content (case-insensitive). Returns matches with context."""
    res = read_file(filepath)
    if res.get("status") != "success":
        return res
    text = res.get("content", "")
    matches = []
    if not keyword:
        return {"status": "error", "error": "empty keyword"}

    pattern = re.compile(re.escape(keyword), re.IGNORECASE)
    for m in pattern.finditer(text):
        start = max(0, m.start() - context)
        end = min(len(text), m.end() + context)
        snippet = text[start:end]
        matches.append({"start": m.start(), "end": m.end(), "snippet": snippet})

    return {"status": "success", "matches": matches, "count": len(matches), "metadata": res.get("metadata")}


def _parse_resume_text(text: str) -> Dict[str, Any]:
    """Lightweight resume parser extracting name, summary, and skills."""
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    name = lines[0] if lines else ""
    # summary = first paragraph (until blank line)
    paragraphs = []
    current = []
    for ln in text.splitlines():
        if ln.strip():
            current.append(ln.strip())
        else:
            if current:
                paragraphs.append(" ".join(current))
                current = []
    if current:
        paragraphs.append(" ".join(current))
    summary = paragraphs[0] if paragraphs else ""

    skills = []
    # find lines like 'Skills: Python, SQL'
    for ln in lines:
        m = re.match(r"skills[:\-]\s*(.+)", ln, re.IGNORECASE)
        if m:
            skills = [s.strip() for s in re.split(r"[,;]", m.group(1)) if s.strip()]
            break

    return {"name": name, "summary": summary, "skills": skills}


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("action", choices=["read", "list", "write", "search"])
    parser.add_argument("path")
    parser.add_argument("-e", "--ext", help="extension for list")
    parser.add_argument("-c", "--content", help="content for write")
    parser.add_argument("-k", "--keyword", help="keyword for search")
    args = parser.parse_args()
    if args.action == "read":
        print(read_file(args.path))
    elif args.action == "list":
        print(list_files(args.path, args.ext))
    elif args.action == "write":
        print(write_file(args.path, args.content or ""))
    elif args.action == "search":
        print(search_in_file(args.path, args.keyword or ""))
